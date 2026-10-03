"""ساخت گزارش HTML (راست‌به‌چپ) و Excel از نامه‌های نخوانده."""
from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .models import LetterSummary

PRIORITY_ORDER = {"فوری": 0, "بالا": 1, "عادی": 2, "کم": 3}

COLUMNS = [
    ("priority", "اولویت"),
    ("referral_type", "نوع ارجاع"),
    ("subject", "موضوع"),
    ("sender", "فرستنده"),
    ("date", "تاریخ"),
    ("letter_number", "شماره نامه"),
    ("summary", "خلاصه"),
    ("action_required", "اقدام لازم"),
    ("deadline", "مهلت"),
]


def sort_rows(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: PRIORITY_ORDER.get(r.get("priority", ""), 9))


def write_reports(rows: list[dict], out_dir: Path) -> tuple[Path, Path]:
    rows = sort_rows(rows)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    html_path = out_dir / f"unread-{stamp}.html"
    xlsx_path = out_dir / f"unread-{stamp}.xlsx"
    html_path.write_text(_html(rows), encoding="utf-8")
    _xlsx(rows, xlsx_path)
    return html_path, xlsx_path


def _xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "نامه‌های نخوانده"
    ws.sheet_view.rightToLeft = True
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(k, "") for k, _ in COLUMNS])
    widths = {"summary": 70, "action_required": 40, "subject": 40}
    for idx, (key, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(key, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def _html(rows: list[dict]) -> str:
    esc = lambda v: html.escape(str(v or ""))  # noqa: E731
    counts: dict[str, int] = {}
    for r in rows:
        counts[r.get("referral_type") or "نامشخص"] = counts.get(r.get("referral_type") or "نامشخص", 0) + 1
    chips = "".join(f'<span class="chip">{esc(k)}: {v}</span>' for k, v in counts.items())
    cards = []
    for r in rows:
        cards.append(
            f"""<article class="card p-{esc(r.get('priority'))}">
  <header><span class="badge">{esc(r.get('priority'))}</span>
  <span class="ref">{esc(r.get('referral_type'))}</span>
  <h2>{esc(r.get('subject'))}</h2></header>
  <p class="meta">{esc(r.get('sender'))} · {esc(r.get('date'))} · شماره {esc(r.get('letter_number'))}</p>
  <p>{esc(r.get('summary'))}</p>
  <p class="action"><b>اقدام لازم:</b> {esc(r.get('action_required'))}{(' · <b>مهلت:</b> ' + esc(r.get('deadline'))) if r.get('deadline') else ''}</p>
</article>"""
        )
    return f"""<!doctype html>
<html lang="fa" dir="rtl"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>نامه‌های نخوانده</title>
<style>
:root {{ --bg:#f6f7f9; --card:#fff; --fg:#1d2330; --muted:#667085; --line:#e4e7ec;
  --urgent:#d92d20; --high:#f79009; --normal:#2e90fa; --low:#98a2b3; }}
@media (prefers-color-scheme: dark) {{ :root {{ --bg:#11151c; --card:#1a1f29; --fg:#e7eaf0; --muted:#98a2b3; --line:#2a303c; }} }}
body {{ margin:0; padding:24px 16px; background:var(--bg); color:var(--fg);
  font-family: Vazirmatn, Tahoma, sans-serif; line-height:1.8; }}
main {{ max-width:880px; margin:auto; }}
h1 {{ font-size:1.4rem; margin:0 0 8px; }}
.chip {{ display:inline-block; margin:0 0 8px 8px; padding:2px 10px; border:1px solid var(--line); border-radius:99px; font-size:.85rem; }}
.card {{ background:var(--card); border:1px solid var(--line); border-right:5px solid var(--normal); border-radius:10px; padding:14px 16px; margin:12px 0; }}
.p-فوری {{ border-right-color:var(--urgent); }} .p-بالا {{ border-right-color:var(--high); }} .p-کم {{ border-right-color:var(--low); }}
.card h2 {{ font-size:1.05rem; margin:6px 0 0; }}
.badge, .ref {{ font-size:.8rem; padding:1px 8px; border-radius:6px; background:var(--line); margin-left:6px; }}
.meta {{ color:var(--muted); font-size:.85rem; margin:4px 0; }}
.action {{ font-size:.92rem; }}
</style></head>
<body><main>
<h1>نامه‌های نخوانده ({len(rows)})</h1>
<div>{chips}</div>
{''.join(cards) or '<p>نامه‌ی نخوانده‌ای نیست.</p>'}
</main></body></html>"""


def load_state(path: Path) -> dict:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}


def save_state(path: Path, state: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def remember(state: dict, summary: LetterSummary) -> None:
    row = summary.to_row()
    row["transcript"] = summary.ocr_text
    row["processed_at"] = datetime.now().isoformat(timespec="seconds")
    state[summary.letter.key] = row
